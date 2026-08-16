"""Post-processing DOCX hasil pdf2docx/LibreOffice untuk Co-Writer.

- Ubah paragraf berukuran besar (>= 14pt) dengan pola bab/sub-bab menjadi
  style Heading 1/2/3 (bukan hanya Normal) supaya struktur tampil di editor.
- Cover page: elemen halaman pertama (sebelum 'BAB 1') dibuat rata tengah,
  TAPI hanya bila paragrafnya belum membawa posisi sendiri (lihat
  ``_punya_indent``) — pdf2docx sudah menaruh geometri sejati tiap baris di
  ``w:ind``, dan merata-tengahkan di atasnya justru menggeser teks.
- Buang skala lebar karakter ``w:w`` (lihat ``_strip_char_scaling``): nilainya
  di luar spesifikasi dan membuat renderer mengabaikan ``w:jc``.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# Versi pipeline impor PDF→DOCX. Dinaikkan setiap kali post-processing di sini
# berubah supaya dokumen yang diimpor dengan pipeline lama dibangun ulang secara
# otomatis (lihat auto-heal di `_prepare_onlyoffice_docx`, co_writer.py). Sidecar
# `uploads/{id}/onlyoffice/pipeline.json` menyimpan versi yang dipakai tiap
# dokumen; bila lebih kecil dari nilai ini DAN pengguna belum menyunting, DOCX
# kerja dibuang lalu dibangun ulang dari PDF asli.
_PIPELINE_IMPOR_VERSI = 2

_BAB = re.compile(r"^(?:bab|chapter)\s+[ivxlcdm\d]+", re.I)
_SUB = re.compile(r"^\d+(\.\d+){1,2}\s+\S")
_FRONT = re.compile(
    r"^(lembar|surat|halaman|kata|daftar|abstrak|abstract|pernyataan|persembahan|"
    r"pedoman|mott|glosarium|biodata|riwayat|lampiran)\b",
    re.I,
)


def _set_heading_style(p, level: int) -> None:
    """Ubah paragraf ke style Heading n; fallback tebal bila style tak ada."""
    try:
        # python-docx versi ini tidak punya p.document; ambil dari part root
        doc = p.part.document
        style = doc.styles[f"Heading {level}"]
        p.style = style
    except Exception:  # noqa: BLE001
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt({1: 16, 2: 14, 3: 12}.get(level, 12))


def _remove_heading_numbering(p) -> None:
    """Hapus auto-numbering (numPr) pada heading supaya tidak dobel nomor."""
    pPr = p._p.get_or_add_pPr()
    for num_pr in pPr.findall(qn("w:numPr")):
        pPr.remove(num_pr)
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    pPr.append(num_pr)


def _center_para(p) -> None:
    """Rata tengah paragraf, sekaligus buang indent yang jadi acuan lamanya.

    Rata tengah menghitung titik tengah dari kotak SETELAH indent dipotong.
    Kalau ``w:ind`` dibiarkan, teks jadi tengah-nya kotak yang tepinya sudah
    tergeser — bukan tengah badan halaman. Semua atribut ``w:ind`` bersifat
    posisi (left/start, right/end, firstLine, hanging), jadi elemennya dibuang
    seluruhnya.
    """
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        for ind in pPr.findall(qn("w:ind")):
            pPr.remove(ind)


# Batas twips di mana indent dianggap sengaja diset, bukan sisa pembulatan.
# w:ind bersatuan twips (1/20 pt), jadi 40 twips = 2pt.
_AMBANG_INDENT_TWIPS = 40

_ATTR_INDENT_HORIZONTAL = ("w:left", "w:start", "w:right", "w:end")


def _punya_indent(p) -> bool:
    """True bila paragraf membawa indent kiri/kanan yang berarti.

    Dipakai untuk mengenali paragraf yang posisinya sudah ditentukan pdf2docx
    (terukur akurat ~0.1pt terhadap PDF asli) supaya tidak ditimpa perataan
    paksa. Nilai dibaca dari atribut XML mentah — satuannya twips, BUKAN EMU
    seperti ``paragraph_format.left_indent`` yang sudah dikonversi python-docx.
    """
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        return False
    for ind in pPr.findall(qn("w:ind")):
        for attr in _ATTR_INDENT_HORIZONTAL:
            raw = ind.get(qn(attr))
            if raw is None:
                continue
            try:
                if abs(int(raw)) > _AMBANG_INDENT_TWIPS:
                    return True
            except ValueError:
                continue
    return False


def postprocess_docx(path: str | Path) -> dict:
    """Perbaiki DOCX hasil konversi: heading & cover. Return statistik."""
    path = Path(path)
    doc = Document(str(path))
    stats = {
        "heading1": 0,
        "heading2": 0,
        "heading3": 0,
        "cover_center": 0,
        "cover_center_skipped": 0,
    }

    paragraphs = doc.paragraphs
    bab_start = len(paragraphs)
    for i, p in enumerate(paragraphs):
        if _BAB.match(p.text.strip()):
            bab_start = i
            break

    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        if not text:
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        size = max(sizes) if sizes else 12.0

        if _BAB.match(text) and size >= 14:
            _set_heading_style(p, 1)
            _remove_heading_numbering(p)
            stats["heading1"] += 1
            continue
        if re.match(r"^[IVXLC]+\.\s+|^\d{1,3}\.\s+\S", text) and size >= 13 and len(text) < 120:
            _set_heading_style(p, 1)
            _remove_heading_numbering(p)
            stats["heading1"] += 1
            continue
        if _SUB.match(text) and size >= 13 and len(text) < 150:
            level = 3 if text.split()[0].count(".") >= 2 else 2
            _set_heading_style(p, level)
            _remove_heading_numbering(p)
            stats[f"heading{level}"] += 1
            continue

        if i < bab_start and size >= 13:
            # Paragraf yang sudah membawa indent posisinya ditentukan konverter
            # dan akurat terhadap PDF asli; merata-tengahkannya menggeser teks
            # (terukur sampai 97pt pada halaman cover). Perataan paksa hanya
            # untuk sumber tanpa data posisi: fallback LibreOffice, markdown,
            # atau HTML.
            if _punya_indent(p):
                stats["cover_center_skipped"] += 1
                continue
            _center_para(p)
            stats["cover_center"] += 1

    doc.save(str(path))

    # Normalisasi spasi antar-run (pdf2docx memecah kata per run tanpa spasi).
    # Dipanggil SETELAH doc.save — _fix_run_spacing menyimpan file sendiri;
    # memanggil sebelum save membuat doc.save menimpa hasil fix dengan objek
    # doc lama yang belum ternormalisasi.
    stats["run_spacing_fixed"] = _fix_run_spacing(path)
    return stats


# ── Lapisan refine pasca pdf2docx (PRD P1) ──────────────────────────────

# Baris daftar dengan dot-leader: "Judul .......... nomor_halaman"
_DOT_LEADER = re.compile(r"^(?P<jdl>.*?)[·.．]{3,}\s*(?P<num>\d{1,4})\s*$")

# Karakter yang TIDAK boleh didahului spasi saat normalisasi run (tanda baca,
# kurung tutup, persen, dst). Spasi tidak pernah hilang sebelum tanda ini.
_NO_SPASI_SEBELUM = set(".,;:!?)]}%‰")

# Karakter yang TIDAK boleh diikuti spasi (kurung buka, dolar, dll).
_NO_SPASI_SESUDAH = set("([{$")

# Selisih ukuran font (pt) yang dianggap "beda ukuran" saat membandingkan dua
# run bertetangga. Longgar sedikit supaya pembulatan half-point tidak terhitung.
_TOLERANSI_UKURAN_PT = 0.05


def _ukuran_run(r) -> float | None:
    """Ukuran efektif run dalam pt (dari w:sz half-point), atau None.

    ``r.font.size`` membaca ``rPr/w:sz`` langsung; bila run tidak menyetelnya,
    rantai style character ditelusuri sebagai fallback.
    """
    if r.font.size is not None:
        return r.font.size.pt
    try:
        style = r.style
        while style is not None:
            if style.font.size is not None:
                return style.font.size.pt
            style = style.base_style
    except Exception:  # noqa: BLE001 — ukuran opsional, jangan gagalkan proses
        return None
    return None


def _digeser_vertikal(r) -> bool:
    """True bila run superskrip/subskrip atau digeser lewat w:position."""
    rPr = r._r.find(qn("w:rPr"))
    if rPr is None:
        return False
    if rPr.find(qn("w:vertAlign")) is not None:
        return True
    pos = rPr.find(qn("w:position"))
    if pos is not None:
        try:
            return int(pos.get(qn("w:val")) or 0) != 0
        except (TypeError, ValueError):
            return False
    return False


def _batas_kata_sungguhan(prev_run, curr_run) -> bool:
    """True bila peralihan dua run memang batas antar-kata.

    pdf2docx tidak cuma memecah per kata: glyph berukuran khusus dipecah jadi
    run sendiri, mis. logo "LaTeX" jadi ``['…LUARAN L', 'A', 'TEX']`` dengan
    ``w:sz`` 20 di antara tetangga 29. Menyisipkan spasi di situ memecah kata
    menjadi "L A TEX". Pada batas kata yang sungguhan kedua run selalu
    seukuran dan sejajar, jadi perbedaan ukuran/pergeseran vertikal dipakai
    sebagai penanda pembeda.
    """
    if _digeser_vertikal(prev_run) or _digeser_vertikal(curr_run):
        return False
    sz_prev = _ukuran_run(prev_run)
    sz_curr = _ukuran_run(curr_run)
    if sz_prev is None or sz_curr is None:
        return True
    return abs(sz_prev - sz_curr) <= _TOLERANSI_UKURAN_PT


def _fix_run_spacing(docx_path: Path) -> int:
    """Normalisasi spasi antar-run paragraf.

    pdf2docx 0.5.13 memecah tiap kata menjadi run terpisah TANPA spasi di
    antaranya (spasi hanya di akhir sebagian run), menghasilkan kata nempel
    ("TugasAkhiriniadalah") saat DOCX dibuka di editor. Sisipkan spasi antar
    run bila: run sebelumnya tidak berakhir spasi, run saat ini tidak berawal
    spasi, dan tidak ada tanda baca yang melarang spasi.
    """
    document = Document(str(docx_path))
    changed = 0
    for p in document.paragraphs:
        runs = p.runs
        for j in range(1, len(runs)):
            prev = runs[j - 1].text or ""
            curr = runs[j].text or ""
            if not prev or not curr:
                continue
            # Lewati bila sudah ada spasi di salah satu sisi.
            if prev[-1].isspace() or curr[0].isspace():
                continue
            # Jangan sisipkan sebelum tanda baca penutup / sesudah pembuka.
            if curr[0] in _NO_SPASI_SEBELUM or prev[-1] in _NO_SPASI_SESUDAH:
                continue
            # Jangan sisipkan di tengah kata yang dipecah per-glyph.
            if not _batas_kata_sungguhan(runs[j - 1], runs[j]):
                continue
            runs[j].text = " " + curr
            changed += 1
    if changed:
        document.save(str(docx_path))
    return changed


def _strip_char_scaling(docx_path: Path) -> int:
    """Buang ``w:rPr/w:w`` (skala lebar karakter) dari seluruh run.

    pdf2docx menulis nilai pecahan seperti ``98.93928396290747`` untuk memaksa
    lebar teks persis sama dengan PDF. Nilai itu di luar spesifikasi —
    ``ST_TextScale`` ECMA-376 hanya menerima persen bulat — dan renderer editor
    (SuperDoc) menangani run ber-``w:w`` di jalur cat terpisah: run dipasang
    ``position: absolute; left: padding-left`` lalu diberi ``scaleX``, sehingga
    ``w:jc`` center/right diabaikan dan teksnya nempel ke tepi indent. Terukur
    pada dokumen 64 halaman: 186 run ber-``w:w`` membuat 37 paragraf melenceng,
    mis. 'LAPORAN TUGAS AKHIR' berakhir di 246.7pt alih-alih 399.9pt dan
    'BANDUNG' mulai di 115.9pt alih-alih di tengah halaman.

    Harga yang dibayar: teks jadi ~1% lebih lebar dari PDF (≈1.8pt untuk judul
    selebar 174pt) — jauh lebih murah daripada geser 150pt+.
    """
    document = Document(str(docx_path))
    changed = 0

    def bersihkan(runs) -> None:
        nonlocal changed
        for r in runs:
            rPr = r._r.find(qn("w:rPr"))
            if rPr is None:
                continue
            for w in rPr.findall(qn("w:w")):
                rPr.remove(w)
                changed += 1

    for p in document.paragraphs:
        bersihkan(p.runs)
    for tbl in document.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    bersihkan(p.runs)
    if changed:
        document.save(str(docx_path))
    return changed


def _fix_bold_from_pdf(docx_path: Path, pdf_path: Path | None) -> int:
    """Koreksi format bold memakai font-flags PDF asli (PyMuPDF bit 0x10).

    pdf2docx mendeteksi bold dari nama font ('Bold' di string), tapi PDF
    dari LaTeX/Word memakai subset font bernama acak ('ABCDEF+Font1'),
    sehingga bold lolos. Flags bit 4 (bold) dari span PDF lebih reliable.
    """
    if not pdf_path or not pdf_path.exists():
        return 0
    import fitz

    bold_words: set[str] = set()
    try:
        with fitz.open(str(pdf_path)) as pdf:
            for page in pdf:
                for block in page.get_text("dict").get("blocks", []):
                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            if span.get("flags", 0) & 0x10:
                                for w in re.split(r"\s+", span.get("text", "").strip()):
                                    if len(w) >= 2:
                                        bold_words.add(w.casefold())
    except Exception:  # noqa: BLE001 — koreksi bold opsional
        return 0

    document = Document(str(docx_path))
    changed = 0
    for p in document.paragraphs:
        for r in p.runs:
            t = r.text.strip()
            if not t or len(t) < 2 or r.font.bold:
                continue
            keys = [w.casefold() for w in re.split(r"\s+", t) if len(w) >= 2]
            if keys and all(k in bold_words for k in keys):
                r.font.bold = True
                changed += 1
    if changed:
        document.save(str(docx_path))
    return changed


def _fix_dot_leader_spacing(docx_path: Path) -> int:
    """Normalisasi line-spacing paragraf berpola dot-leader (PRD P1 bug 2).

    pdf2docx salah menghitung tinggi baris pada baris 'Judul .... nomor'
    (multi-run dgn x jauh), baris berikutnya tampak naik. Samakan spacing
    dgn paragraf body sekitarnya.
    """
    document = Document(str(docx_path))
    baseline = None
    for p in document.paragraphs:
        t = p.text.strip()
        if t and not _DOT_LEADER.search(t):
            ls = p.paragraph_format.line_spacing
            if ls:
                baseline = ls
                break
    changed = 0
    for p in document.paragraphs:
        if not _DOT_LEADER.search(p.text):
            continue
        pf = p.paragraph_format
        target = baseline or 1.5
        if pf.line_spacing != target:
            pf.line_spacing = target
            changed += 1
    if changed:
        document.save(str(docx_path))
    return changed


def post_process_converted_docx(
    docx_path: str | Path,
    source_pdf_path: str | Path | None = None,
) -> dict:
    """Refine akhir DOCX hasil pdf2docx: skala karakter, bold, spacing dot-leader."""
    docx_path = Path(docx_path)
    stats = {"char_scaling_stripped": 0, "bold_fixed_runs": 0, "dot_leader_fixed": 0}
    stats["char_scaling_stripped"] = _strip_char_scaling(docx_path)
    if source_pdf_path is not None:
        stats["bold_fixed_runs"] = _fix_bold_from_pdf(docx_path, Path(source_pdf_path))
    stats["dot_leader_fixed"] = _fix_dot_leader_spacing(docx_path)
    return stats