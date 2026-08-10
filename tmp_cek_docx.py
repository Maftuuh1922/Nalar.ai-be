"""Cek DOCX yang dibuka di OnlyOffice: marker bocor, H1, struktur."""

import os
import docx

path = os.path.join(os.environ.get("TEMP", ""), "diag_onlyoffice.docx")
d = docx.Document(path)

bocor = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if "<center>" in t or "</center>" in t or "<newpage>" in t or t.startswith("\\") or "## " in t or "**" in t:
        bocor.append((i, p.style.name, t[:100]))
print("== Marker/LaTeX/penanda bocor ke DOCX:", len(bocor))
for b in bocor[:25]:
    print(b)

h = [p for p in d.paragraphs if p.style.name.startswith("Heading")]
print("\nHeading 1:", [p.text[:60] for p in d.paragraphs if p.style.name == "Heading 1"])
print("\nungguh heading 2 pertama:", h[12].text, "|", h[13].text, "|", h[14].text)