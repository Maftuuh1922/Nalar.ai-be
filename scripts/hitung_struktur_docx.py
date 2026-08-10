"""Hitung struktur asli sebuah DOCX — pembanding untuk mengukur rugi impor.

Angka di sini adalah "yang diinputkan". Kalau .tex hasil impor tidak
mendekatinya, kerusakan terjadi di pengekstrak, bukan di pengekspor.
"""

from __future__ import annotations

import sys
import zipfile
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn


def main() -> int:
    jalur = Path(sys.argv[1])
    d = docx.Document(str(jalur))

    gaya = Counter(p.style.name for p in d.paragraphs)
    tebal = sum(
        1 for p in d.paragraphs for r in p.runs if r.bold and (r.text or "").strip()
    )
    miring = sum(
        1 for p in d.paragraphs for r in p.runs if r.italic and (r.text or "").strip()
    )
    # numPr = paragraf yang benar-benar berada di dalam daftar bernomor/butir.
    daftar = sum(1 for p in d.paragraphs if p._p.find(qn("w:pPr")) is not None
                 and p._p.find(qn("w:pPr")).find(qn("w:numPr")) is not None)
    jeda = sum(
        1
        for p in d.paragraphs
        for br in p._p.iter(qn("w:br"))
        if br.get(qn("w:type")) == "page"
    )
    with zipfile.ZipFile(jalur) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]

    print(f"paragraf          {len(d.paragraphs)}")
    print(f"tabel             {len(d.tables)}")
    print(f"gambar tertanam   {len(media)}")
    print(f"run tebal         {tebal}")
    print(f"run miring        {miring}")
    print(f"paragraf daftar   {daftar}")
    print(f"jeda halaman      {jeda}")
    print("\ngaya paragraf terbanyak:")
    for nama, n in gaya.most_common(14):
        print(f"  {n:>5}  {nama}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
