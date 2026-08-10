"""Simulasi jalur DOCX yang dibuka di OnlyOffice dari PDF hasil impor."""
import sys, os, tempfile

sys.path.insert(0, r"C:\Users\Administrator\Documents\project ta\Nalar.ai-be")

SUMBER = r"C:\Users\Administrator\Documents\project ta\Laporan_Tugas_Akhir_Nalar_AI.pdf"
IMAGES = r"C:\Users\ADMINI~1\AppData\Local\Temp\opencode\diag_import\images"

from app.services.doc_import import pdf_to_markdown, _rapikan
from app.services.doc_ast import markdown_to_ast, ast_to_markdown
from app.services.docx_template_exporter import markdown_to_docx_template

md = pdf_to_markdown(open(SUMBER, "rb").read(), IMAGES, "http://localhost:8087/uploads/d/images", "d")
bersih = _rapikan(md, preserve_content=False)
ast = markdown_to_ast(bersih, title="Laporan TA")
ast_md = ast_to_markdown(ast)

out = os.path.join(tempfile.gettempdir(), "diag_onlyoffice.docx")
markdown_to_docx_template(ast_md, out, None)
print("DOCX:", out, os.path.getsize(out)//1024, "KB")

import docx
d = docx.Document(out)
styles = {}
gaya = 0
for p in d.paragraphs:
    s = (p.style.name or "")
    if s.startswith("Heading"):
        gaya += 1
        if gaya <= 12:
            print(f"H{gaya}: [{s}] {p.text[:80]}")
    styles.setdefault(s, 0)
    styles[s] += 1
print("\nJumlah Heading:", gaya)
print("Total paragraf:", len(d.paragraphs))
print("Gaya:", styles)